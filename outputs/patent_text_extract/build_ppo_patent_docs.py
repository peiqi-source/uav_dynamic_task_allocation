from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(r"F:\项目\专利\uav_dynamic_task_allocation")
OUT_DIR = ROOT / "专利" / "PPO" / "基于PPO的无人机集群动态目标分群专利"
FIG_DIR = OUT_DIR / "附图"
DEMO_IMAGE = ROOT / "outputs" / "evaluation" / "figures" / "ppo_dynamic_added_targets_before_after.png"


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = Pt(22)
    normal.paragraph_format.space_after = Pt(0)


def set_font(run, name: str = "宋体", size: int = 12, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc: Document, text: str, indent: bool = True, align=None, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = Pt(22)
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_font(run, bold=bold)


def add_title(doc: Document, text: str) -> None:
    add_para(doc, text, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)


def add_heading(doc: Document, text: str) -> None:
    add_para(doc, text, indent=False, bold=True)


def add_formula(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = Pt(18)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_font(run, name="Times New Roman", size=11)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line([start, end], fill=(30, 30, 30), width=3)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        points = [(x2, y2), (x2 - 12, y2 - 7), (x2 - 12, y2 + 7)]
    else:
        points = [(x2, y2), (x2 + 12, y2 - 7), (x2 + 12, y2 + 7)]
    draw.polygon(points, fill=(30, 30, 30))


def draw_flow(path: Path, title: str, boxes: list[str], rows: int = 1) -> None:
    width, height = 1500, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(34, bold=True)
    text_font = load_font(26)
    small_font = load_font(22)
    draw.text((width // 2, 45), title, anchor="mm", fill=(0, 0, 0), font=title_font)
    if rows == 1:
        box_w, box_h = 190, 95
        gap = 38
        y = 320
        total = len(boxes) * box_w + (len(boxes) - 1) * gap
        x = (width - total) // 2
        centers = []
        for box in boxes:
            rect = (x, y, x + box_w, y + box_h)
            draw.rounded_rectangle(rect, radius=14, outline=(20, 20, 20), width=3, fill=(244, 247, 250))
            draw.multiline_text((x + box_w // 2, y + box_h // 2), box, anchor="mm", fill=(0, 0, 0), font=text_font, align="center", spacing=5)
            centers.append((x + box_w, y + box_h // 2, x + box_w + gap, y + box_h // 2))
            x += box_w + gap
        for x1, y1, x2, y2 in centers[:-1]:
            arrow(draw, (x1 + 5, y1), (x2 - 5, y2))
    else:
        box_w, box_h = 300, 90
        coords = [(190, 185), (610, 185), (1030, 185), (190, 430), (610, 430), (1030, 430)]
        for idx, box in enumerate(boxes):
            x, y = coords[idx]
            rect = (x, y, x + box_w, y + box_h)
            draw.rounded_rectangle(rect, radius=14, outline=(20, 20, 20), width=3, fill=(244, 247, 250))
            draw.multiline_text((x + box_w // 2, y + box_h // 2), box, anchor="mm", fill=(0, 0, 0), font=small_font, align="center", spacing=5)
        for start, end in [((490, 230), (610, 230)), ((910, 230), (1030, 230)), ((1180, 275), (340, 430)), ((490, 475), (610, 475)), ((910, 475), (1030, 475))]:
            arrow(draw, start, end)
    image.save(path)


def build_figures() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    draw_flow(FIG_DIR / "图1.png", "基于PPO的无人机集群动态目标分群总体流程", ["摧毁目标集", "PSO初始分群", "动态事件", "PPO决策", "簇标签更新", "目标簇输出"])
    draw_flow(FIG_DIR / "图2.png", "摧毁目标集与PSO初始分群构建流程", ["目标状态向量", "特征归一化", "粒子群搜索", "初始标签", "簇中心"], rows=1)
    draw_flow(FIG_DIR / "图3.png", "动态目标事件触发与状态更新流程", ["新增目标", "目标消失", "属性变化", "当前目标集合", "簇级统计更新"], rows=1)
    draw_flow(FIG_DIR / "图4.png", "PPO状态、动作和奖励构建示意", ["簇级状态", "动作掩码", "目标移动动作", "分群得分", "奖励反馈", "下一状态"], rows=2)
    draw_flow(FIG_DIR / "图5.png", "PPO动态重分群训练流程", ["真实目标数据", "动态事件场景", "环境交互", "轨迹样本", "PPO裁剪更新", "策略模型"], rows=1)
    if DEMO_IMAGE.exists():
        Image.open(DEMO_IMAGE).convert("RGB").save(FIG_DIR / "图6.png")
    else:
        draw_flow(FIG_DIR / "图6.png", "新增目标事件下重分群前后对比", ["新增目标", "重分群前", "PPO推理", "重分群后"], rows=1)
    draw_flow(FIG_DIR / "图7.png", "动态重分群接入资源分配与打击规划", ["目标簇集合", "资源分配", "完成率评估", "打击次序规划", "任务反馈"], rows=1)


SPEC_PARAGRAPHS = {
    "技术领域": [
        "本发明属于无人机集群任务规划、动态目标分群、强化学习决策和智能无人作战系统技术领域，具体涉及一种基于近端策略优化的无人机集群动态目标分群方法。",
    ],
    "背景技术": [
        "现代战争朝着智能化、无人化和体系化方向发展，无人机集群凭借部署灵活、协同能力强、任务适应性好等特点，已成为复杂战场环境下执行侦察、监视、打击和协同支援任务的重要装备形式。在对地目标任务中，系统通常先确定摧毁目标集，再对摧毁目标集进行区域目标划分，以便不同无人机子群围绕相对集中的目标区域开展资源分配和打击规划。",
        "静态区域目标划分可采用粒子群优化等智能优化方法，根据目标位置、目标价值和目标防御能力搜索目标簇中心，得到较为紧凑的初始目标区域。该类方法适用于任务开始前的全局初始划分，能够为无人机资源配置提供基础。",
        "但是，在开放战场环境中，目标态势会随任务执行过程发生变化。例如，侦察系统发现新的敌方目标，原有目标被其他作战力量摧毁、撤离或失去侦测信号，目标重要程度和防御能力随战场态势发生变化。上述动态事件会改变原摧毁目标集的空间分布和任务负载，使原初始分群不再满足任务均衡和区域紧凑要求。",
        "若每次动态事件发生后均完全重新聚类，容易破坏既有任务结构，使已分配无人机频繁改变任务区域，造成任务连续性差、协同通信开销增大以及实时计算压力增高。固定规则式局部调整方法又难以同时兼顾空间紧凑性、目标数量均衡性、价值防御负载均衡性以及非法目标迁移动作约束。",
        "因此，有必要提出一种在摧毁目标集已经确定、静态初始区域划分已经完成的基础上，根据动态目标事件对当前目标分群状态进行连续决策调整的方法，使其能够在目标新增、目标消失和目标属性变化后，快速输出稳定、紧凑且任务负载均衡的动态目标簇集合。",
    ],
    "发明内容": [
        "为了克服现有技术在动态目标事件下频繁全局重新分群、任务结构连续性不足和任务负载均衡表达不充分的问题，本发明提供了一种基于近端策略优化的无人机集群动态目标分群方法。该方法首先基于已经确定的摧毁目标集采用粒子群优化生成初始目标簇，然后在动态目标事件发生后，将当前目标簇状态输入近端策略优化智能体，由智能体输出目标移动动作，实现对原有分群结果的动态调整。",
        "本发明解决其技术问题所采用的技术方案如下：",
        "步骤1：获取已经确定的摧毁目标集，构造目标状态向量；",
        "步骤2：对目标状态向量进行归一化处理，并采用粒子群优化方法生成初始目标簇标签和初始目标簇中心；",
        "步骤3：检测动态目标事件，所述动态目标事件包括目标新增、目标消失、目标属性变化和混合事件；",
        "步骤4：根据动态目标事件更新当前摧毁目标集合，并将更新后的目标集合与原始分群标签合并形成当前分群状态；",
        "步骤5：构建近端策略优化智能体的簇级状态空间；",
        "步骤6：构建目标移动动作空间，动作表示为将源目标簇中的一个目标移动至目标目标簇；",
        "步骤7：构建综合区域紧凑性、目标数量均衡性、任务负载均衡性和动作惩罚的奖励函数；",
        "步骤8：采用训练后的近端策略优化策略网络对当前分群状态进行推理，输出目标移动动作；",
        "步骤9：根据目标移动动作更新目标簇标签和目标簇中心；",
        "步骤10：输出动态重分群后的目标簇集合，供无人机资源分配、任务完成率评估和打击次序规划使用。",
        "进一步地，所述步骤1具体为：设摧毁目标集合为D={T_i|i=1,2,...,N_d}，其中N_d为摧毁目标数量，第i个目标T_i的状态向量表示为：",
        "其中，x_i和y_i为目标坐标，s_i为目标重要度，d_i为目标防御能力，q_i为目标综合评分，type_i为目标类型。",
        "进一步地，所述步骤2具体为：对目标状态向量的第r维特征进行归一化处理，得到归一化特征；以目标簇中心为粒子位置，构建粒子群优化适应度函数，搜索使区域紧凑性、数量均衡性和任务负载均衡性综合最优的初始分群结果。",
        "进一步地，所述步骤5具体为：对于第k个目标簇G_k，计算簇级状态向量g_k。所述簇级状态向量包括目标数量、簇中心坐标、簇内平均距离、坐标标准差、价值总和、防御总和、任务负载和负载比例。所有目标簇的簇级状态向量按固定顺序拼接形成近端策略优化智能体的观测状态。",
        "进一步地，所述步骤6具体为：目标移动动作表示为a_t=(k_s,m,k_d)，其中k_s为源目标簇编号，m为源目标簇内目标序号，k_d为目标目标簇编号。若k_s与k_d相同、源目标簇为空或m超过源目标簇目标数量，则判定为非法动作，并给予惩罚而不中断流程。",
        "进一步地，所述步骤7具体为：以执行动作前后的分群综合得分差作为基础奖励，并叠加改善奖励、移动代价、空簇惩罚和非法动作惩罚。所述分群综合得分由目标数量均衡得分、区域紧凑性得分和任务负载均衡得分构成。",
        "进一步地，所述近端策略优化训练过程不使用人工分群标签，而是基于真实战场目标态势数据生成动态事件场景，通过目标分群环境交互采集状态、动作、奖励和下一状态序列，并使用近端策略优化裁剪目标函数更新策略网络和价值网络。",
        "本发明的有益效果是：能够在目标新增、目标消失和目标属性变化后快速调整目标区域；能够以初始粒子群分群结果为基础进行局部连续调整，避免完全重新分群造成任务震荡；能够同时兼顾区域紧凑性、目标数量均衡性和任务负载均衡性；能够通过非法动作惩罚和终止条件提高动态调整稳定性；能够与摧毁目标集决策、资源分配、蒙特卡洛资源评估和打击次序规划形成完整任务链。",
    ],
    "附图说明": [
        "图1为本发明基于PPO的无人机集群动态目标分群方法总体流程图。",
        "图2为本发明摧毁目标集与PSO初始分群构建流程图。",
        "图3为本发明动态目标事件触发与状态更新流程图。",
        "图4为本发明PPO智能体状态、动作和奖励构建示意图。",
        "图5为本发明PPO动态重分群训练流程图。",
        "图6为本发明新增目标事件下重分群前后对比示意图。",
        "图7为本发明动态重分群结果接入无人机资源分配与打击规划流程图。",
    ],
    "具体实施方式": [
        "以下结合附图对本发明作进一步详细说明。应当理解，以下实施方式用于说明本发明而非限定本发明的保护范围。在不背离本发明构思的情况下，目标簇数量、奖励权重、训练回合数量、动态事件概率和推理步数均可根据任务规模进行调整。",
        "步骤S1：获取战场目标数据与摧毁目标集。所述摧毁目标集由前置目标筛选模型、任务规则或态势分析结果给出，本发明不重复执行摧毁目标选择，而是以已经确定的摧毁目标集作为动态重分群输入。",
        "步骤S2：基于目标位置、重要度、防御值和综合评分构造目标状态向量，并进行归一化处理。若某一特征的最大值与最小值相同，则将其归一化分母设置为预设常数，避免除零。",
        "步骤S3：采用PSO对摧毁目标集进行初始区域划分。设目标簇数量为K，初始分群标签为l_i，目标簇中心为mu_k，目标簇任务负载为W_k，则PSO适应度函数可以设置为：",
        "式中，C为区域紧凑项，B_n为目标数量均衡项，B_w为任务负载均衡项，alpha、beta和gamma为权重系数。通过最小化F得到初始目标簇标签L_0和初始目标簇中心Mu_0。",
        "步骤S4：检测动态目标事件。当出现新增目标时，将新增目标加入当前目标集合；当出现目标消失时，将消失目标从当前目标集合中删除；当出现目标属性变化时，更新目标重要度、防御能力和综合评分；当多种事件同时出现时，按目标编号保持目标身份一致并统一更新当前目标集合。",
        "步骤S5：将动态事件后的目标集合与原始分群结果合并。对新增目标，可先按距离最近的初始簇中心确定临时标签；对消失目标，删除其标签并更新所在簇统计信息；对属性变化目标，保留原标签并更新簇级价值、防御和负载信息。",
        "步骤S6：构建PPO状态空间。对于第k个目标簇，簇级状态向量包括n_k、c_k^x、c_k^y、r_k、sigma_k、V_k、D_k、W_k和rho_k，其中n_k为簇内目标数量，c_k^x和c_k^y为簇中心坐标，r_k为簇内平均距离，sigma_k为坐标离散程度，V_k为价值总和，D_k为防御总和，W_k为任务负载，rho_k为负载比例。",
        "步骤S7：构建PPO动作空间。动作a_t=(k_s,m,k_d)表示将源目标簇G_{k_s}中第m个目标移动到目标目标簇G_{k_d}。该动作空间将动态重分群转化为若干连续目标迁移操作，避免每次动态事件后均完全重新聚类。",
        "步骤S8：构建奖励函数。设Q_t为动作执行前的分群综合得分，Q_{t+1}为动作执行后的分群综合得分，I_+表示动作是否带来得分改善，I_illegal表示动作是否合法，则单步奖励表示为：",
        "式中，eta为改善奖励系数，c_m为移动代价，c_i为非法动作惩罚系数，clip表示奖励裁剪函数。通过奖励裁剪可降低训练过程中的大幅波动，提高收敛性展示效果。",
        "步骤S9：训练PPO智能体。以真实目标态势数据作为基础目标分布，按照无扰动、目标删除、目标新增、目标属性变化和混合事件生成多样化训练回合。每个回合中先由PSO得到初始分群，再由智能体与分群环境交互，采集(s_t,a_t,r_t,s_{t+1})序列。",
        "步骤S10：采用近端策略优化裁剪目标函数更新策略网络。策略比率和裁剪目标函数表示为：",
        "式中，pi_theta为当前策略，pi_theta_old为旧策略，A_t为优势函数，epsilon为裁剪系数。通过限制策略更新幅度，使动态分群策略在训练过程中保持稳定。",
        "步骤S11：推理阶段加载训练后的策略网络，以PSO初始分群为起点，根据动作掩码选择合法目标移动动作。推理过程在达到最大步数、无合法动作或连续若干步无得分改善时结束。",
        "步骤S12：输出动态重分群后的目标簇集合。所述目标簇集合包括目标簇编号、簇内目标编号、目标簇中心、目标数量、价值总和、防御总和、任务负载和分群评价指标，并输入后续无人机资源分配、蒙特卡洛任务完成率评估和打击次序规划。",
        "在一组实施例中，基于同一战场目标态势数据分别进行PSO初始分群和PPO动态重分群评价。结果表明，PPO动态重分群在保持空簇数量为零的情况下，能够提高分群综合得分、目标数量均衡性和任务负载均衡性，说明训练后的策略能够在动态事件场景下对原有分群结构进行有效调整。",
    ],
}


FORMULAS = {
    "state": "v_i=[x_i, y_i, type_i, s_i, d_i, q_i]^T",
    "norm": "hat(v)_{i,r}=(v_{i,r}-min(v_r))/(max(v_r)-min(v_r)+epsilon)",
    "pso": "F=alpha*C+beta*B_n+gamma*B_w",
    "compact": "C=(1/N_d) * sum_{i=1}^{N_d} ||p_i-mu_{l_i}||_2",
    "balance": "B_n=std(n_1,n_2,...,n_K),  W_k=sum_{i:l_i=k}(s_i+d_i),  B_w=std(W_1,W_2,...,W_K)",
    "cluster_state": "g_k=[n_k,c_k^x,c_k^y,r_k,sigma_k,V_k,D_k,W_k,rho_k]",
    "score": "Q=lambda_1/(1+std(n))+lambda_2/(1+bar(r))+lambda_3/(1+std(W))-lambda_4*E",
    "reward": "r_t=clip(Q_{t+1}-Q_t+eta*I_+-c_m-c_i*I_illegal, r_min, r_max)",
    "ratio": "rho_t(theta)=pi_theta(a_t|s_t)/pi_theta_old(a_t|s_t)",
    "clip": "L^CLIP(theta)=E[min(rho_t(theta)A_t, clip(rho_t(theta),1-epsilon,1+epsilon)A_t)]",
}


def build_specification() -> None:
    doc = Document()
    configure_document(doc)
    add_title(doc, "基于近端策略优化的无人机集群动态目标分群方法")
    for heading in ["技术领域", "背景技术", "发明内容"]:
        add_heading(doc, heading)
        for text in SPEC_PARAGRAPHS[heading]:
            add_para(doc, text)
            if "状态向量表示为" in text:
                add_formula(doc, FORMULAS["state"])
            if "归一化特征" in text:
                add_formula(doc, FORMULAS["norm"])
            if "簇级状态向量g_k" in text:
                add_formula(doc, FORMULAS["cluster_state"])
            if "分群综合得分" in text and "所述分群综合得分由" in text:
                add_formula(doc, FORMULAS["score"])
    add_heading(doc, "附图说明")
    for text in SPEC_PARAGRAPHS["附图说明"]:
        add_para(doc, text)
    add_heading(doc, "具体实施方式")
    for text in SPEC_PARAGRAPHS["具体实施方式"]:
        add_para(doc, text)
        if "PSO适应度函数" in text:
            add_formula(doc, FORMULAS["pso"])
            add_formula(doc, FORMULAS["compact"])
            add_formula(doc, FORMULAS["balance"])
        if "簇级状态向量包括" in text:
            add_formula(doc, FORMULAS["cluster_state"])
        if "单步奖励表示为" in text:
            add_formula(doc, FORMULAS["reward"])
        if "裁剪目标函数表示为" in text:
            add_formula(doc, FORMULAS["ratio"])
            add_formula(doc, FORMULAS["clip"])
    doc.save(OUT_DIR / "说明书.docx")


def build_claims() -> None:
    doc = Document()
    configure_document(doc)
    claims = [
        "1、一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，包括如下步骤：",
        "步骤1：获取已经确定的摧毁目标集，构造目标状态向量；",
        "步骤2：采用粒子群优化方法对所述摧毁目标集进行初始区域划分，得到初始目标簇标签和初始目标簇中心；",
        "步骤3：检测动态目标事件，并根据所述动态目标事件更新当前目标集合；",
        "步骤4：根据当前目标集合和初始目标簇标签构建近端策略优化智能体的簇级状态空间；",
        "步骤5：构建目标移动动作空间，所述目标移动动作空间表示将源目标簇中的目标移动至目标目标簇；",
        "步骤6：构建包含区域紧凑性、目标数量均衡性、任务负载均衡性、空簇惩罚、非法动作惩罚和得分改善量的奖励函数；",
        "步骤7：利用训练后的近端策略优化策略网络输出目标移动动作，并根据所述目标移动动作更新目标簇标签和目标簇中心；",
        "步骤8：输出动态重分群后的目标簇集合。",
        "2、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述步骤1具体为：设摧毁目标集合为D={T_i|i=1,2,...,N_d}，第i个目标的状态向量为：",
        FORMULAS["state"],
        "其中，x_i和y_i为目标坐标，type_i为目标类型，s_i为目标重要度，d_i为目标防御能力，q_i为目标综合评分。",
        "3、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述步骤2具体为：对目标状态向量进行归一化处理：",
        FORMULAS["norm"],
        "以目标簇中心作为粒子位置，基于如下适应度函数搜索初始目标簇标签：",
        FORMULAS["pso"],
        "其中，C为区域紧凑项，B_n为目标数量均衡项，B_w为任务负载均衡项，alpha、beta和gamma为权重系数。",
        "4、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述动态目标事件包括目标新增事件、目标消失事件、目标属性变化事件和混合事件。",
        "5、根据权利要求4所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，当所述动态目标事件为目标新增事件时，将新增目标接入当前目标集合，并根据新增目标与初始目标簇中心之间的距离确定临时目标簇标签，或通过近端策略优化策略网络确定其目标簇标签。",
        "6、根据权利要求4所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，当所述动态目标事件为目标消失事件时，从当前目标集合中删除消失目标，并更新对应目标簇的目标数量、目标簇中心、区域紧凑性和任务负载。",
        "7、根据权利要求4所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，当所述动态目标事件为目标属性变化事件时，更新对应目标的重要度、防御能力、综合评分和任务负载，并基于更新后的目标属性重新计算当前分群状态。",
        "8、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述簇级状态空间中第k个目标簇的状态向量为：",
        FORMULAS["cluster_state"],
        "其中，n_k为簇内目标数量，c_k^x和c_k^y为簇中心坐标，r_k为簇内平均距离，sigma_k为坐标离散程度，V_k为价值总和，D_k为防御总和，W_k为任务负载，rho_k为负载比例。",
        "9、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述目标移动动作表示为a_t=(k_s,m,k_d)，其中k_s为源目标簇编号，m为源目标簇中的目标序号，k_d为目标目标簇编号。",
        "10、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述奖励函数为：",
        FORMULAS["reward"],
        "其中，Q_t为动作执行前的分群综合得分，Q_{t+1}为动作执行后的分群综合得分，I_+表示动作是否使分群得分改善，I_illegal表示动作是否合法，c_m为移动代价，c_i为非法动作惩罚系数。",
        "11、根据权利要求10所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述分群综合得分Q为：",
        FORMULAS["score"],
        "其中，std(n)为各目标簇目标数量标准差，bar(r)为簇内平均距离，std(W)为各目标簇任务负载标准差，E为空簇数量。",
        "12、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述近端策略优化策略网络的训练不使用人工分群标签，而是基于真实目标态势数据和动态事件扰动生成训练回合，通过目标分群环境交互采集状态、动作、奖励和下一状态序列。",
        "13、根据权利要求12所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述近端策略优化策略网络采用裁剪目标函数更新：",
        FORMULAS["clip"],
        "其中，rho_t(theta)为当前策略与旧策略在状态s_t下选择动作a_t的概率比，A_t为优势函数，epsilon为裁剪系数。",
        "14、根据权利要求1所述的一种基于近端策略优化的无人机集群动态目标分群方法，其特征在于，所述动态重分群后的目标簇集合被输出至无人机资源分配模块、任务完成率评估模块或打击次序规划模块。",
        "15、一种基于近端策略优化的无人机集群动态目标分群系统，其特征在于，包括目标数据获取模块、初始分群模块、动态事件处理模块、状态构建模块、动作决策模块、分群更新模块和结果输出模块，各模块用于执行权利要求1至14任一项所述的方法。",
        "16、一种电子设备，包括处理器和存储器，所述存储器存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至14任一项所述的方法。",
        "17、一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至14任一项所述的方法。",
    ]
    for claim in claims:
        add_para(doc, claim, indent=False)
    doc.save(OUT_DIR / "权利要求书.docx")


def build_abstract() -> None:
    doc = Document()
    configure_document(doc)
    abstract = (
        "本发明公开了一种基于近端策略优化的无人机集群动态目标分群方法，属于无人机集群任务规划与强化学习决策领域。"
        "该方法在摧毁目标集已经确定并完成PSO初始区域划分后，针对目标新增、目标消失和目标属性变化等动态事件，构建簇级状态空间和目标移动动作空间；"
        "簇级状态包括目标数量、簇中心、簇内平均距离、坐标离散程度、价值总和、防御总和和任务负载，目标移动动作表示将源目标簇中的目标迁移至目标目标簇；"
        "奖励函数综合区域紧凑性、目标数量均衡性、任务负载均衡性、空簇惩罚、非法动作惩罚和分群得分改善量，并通过近端策略优化裁剪目标函数训练策略网络。"
        "推理时以PSO初始分群为起点，由训练后的策略网络连续输出目标移动动作，更新目标簇标签和簇中心，形成动态重分群后的目标簇集合。"
        "本发明能够提高动态战场中无人机集群目标区域划分的连续性、均衡性和适应能力，为后续资源分配、任务完成率评估和打击次序规划提供可靠输入。"
    )
    add_para(doc, abstract, indent=True)
    doc.save(OUT_DIR / "说明书摘要.docx")


def build_fig_doc() -> None:
    doc = Document()
    configure_document(doc)
    captions = [
        "图1 本发明基于PPO的无人机集群动态目标分群方法总体流程图",
        "图2 本发明摧毁目标集与PSO初始分群构建流程图",
        "图3 本发明动态目标事件触发与状态更新流程图",
        "图4 本发明PPO智能体状态、动作和奖励构建示意图",
        "图5 本发明PPO动态重分群训练流程图",
        "图6 本发明新增目标事件下重分群前后对比示意图",
        "图7 本发明动态重分群结果接入无人机资源分配与打击规划流程图",
    ]
    for idx, caption in enumerate(captions, start=1):
        add_para(doc, caption, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(FIG_DIR / f"图{idx}.png"), width=Inches(6.2))
    doc.save(OUT_DIR / "说明书附图.docx")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_figures()
    build_specification()
    build_claims()
    build_abstract()
    build_fig_doc()
    old_combined = OUT_DIR / "摘要及附图.docx"
    if old_combined.exists():
        old_combined.unlink()
    print(f"wrote four patent files to {OUT_DIR}")


if __name__ == "__main__":
    main()
