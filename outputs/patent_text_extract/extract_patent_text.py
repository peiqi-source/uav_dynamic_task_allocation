from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
PATHS = [
    ROOT / "专利/RF摧毁目标集智能决策专利/说明书.docx",
    ROOT / "专利/RF摧毁目标集智能决策专利/权利要求书.docx",
    ROOT / "专利/RF摧毁目标集智能决策专利/说明书摘要.docx",
    ROOT / "专利/RF摧毁目标集智能决策专利/RF摧毁目标集智能决策方法-说明书附图.docx",
    ROOT / "专利/基于BAB-枚举组合法的无人机多目标攻击航迹规划求解方法/说明书.docx",
    ROOT / "专利/基于BAB-枚举组合法的无人机多目标攻击航迹规划求解方法/权利要求书.docx",
    ROOT / "专利/基于BAB-枚举组合法的无人机多目标攻击航迹规划求解方法/说明书摘要.docx",
    ROOT / "专利/报告.docx",
]


def extract(path: Path) -> None:
    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for ti, table in enumerate(doc.tables):
        lines.append(f"[[TABLE {ti + 1}]]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    target = path.parent.name + "_" + path.stem + ".txt"
    out_path = Path(__file__).resolve().parent / target
    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(path, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "chars", sum(len(x) for x in lines))


def main() -> None:
    for path in PATHS:
        extract(path)


if __name__ == "__main__":
    main()
